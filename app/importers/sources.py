"""Источники, кроме бумажного чека: JSON от ОФД и таблицы выгрузок.

PDF и текст разбирает ofd_pdf. Здесь два формата, которые встречаются чаще всего:

- **JSON чека** — то, что отдаёт «Мои чеки онлайн» ФНС и операторы фискальных данных.
  Суммы там в копейках целыми числами, поэтому целое делим на 100, а дробное считаем
  уже рублями. Структура у разных ОФД вложена по-разному, поэтому ищем первый объект,
  в котором есть список items.
- **Таблица** — CSV или XLSX из личного кабинета магазина либо собранная руками.
  Колонки узнаются по названию, лишние игнорируются.
"""
from __future__ import annotations

import csv
import json
import os
import re
from datetime import datetime
from typing import Any

from app.importers.ofd_pdf import Receipt, ReceiptRow

# как называются нужные колонки в выгрузках
COLUMNS = {
    "name": ("наименование", "название", "товар", "позиция", "продукт", "name", "item"),
    "qty": ("количество", "кол-во", "колво", "кол", "quantity", "qty", "amount"),
    "price": ("цена", "цена за единицу", "цена, ₽", "price", "unit_price"),
    "total": ("сумма", "стоимость", "итог", "total", "sum"),
    "date": ("дата", "date", "дата покупки"),
    "store": ("магазин", "store", "точка", "продавец"),
}


def _norm(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _pick_column(header: list[str], kind: str) -> int | None:
    names = COLUMNS[kind]
    normalized = [_norm(h) for h in header]
    for i, cell in enumerate(normalized):
        if cell in names:
            return i
    for i, cell in enumerate(normalized):           # частичное совпадение
        if any(cell.startswith(n) or n in cell for n in names):
            return i
    return None


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d,.\-]", "", str(value)).replace(",", ".")
    if cleaned.count(".") > 1:                       # 1.234.56 -> 1234.56
        head, _, tail = cleaned.rpartition(".")
        cleaned = head.replace(".", "") + "." + tail
    try:
        return float(cleaned)
    except ValueError:
        return None


def _money(value: Any) -> float | None:
    """Целое — копейки (так отдаёт ФНС), дробное — уже рубли."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return round(value / 100, 2)
    number = _number(value)
    return round(number, 2) if number is not None else None


def _date(value: Any) -> str:
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S",
                    "%Y-%m-%d", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
        try:
            return datetime.strptime(text[:len(datetime.now().strftime(pattern))], pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", text) or re.search(r"(\d{2})\.(\d{2})\.(\d{4})", text)
    if match:
        parts = match.groups()
        return f"{parts[0]}-{parts[1]}-{parts[2]}" if len(parts[0]) == 4 else f"{parts[2]}-{parts[1]}-{parts[0]}"
    return datetime.now().strftime("%Y-%m-%d")


# ---------- JSON чека ----------
_EAN = re.compile(r"(?<!\d)(\d{12,14}|\d{8})(?!\d)")     # длинный код важнее короткого


def _barcode(item: dict) -> str | None:
    """Штрихкод позиции чека, если ОФД его передал.

    Формат у операторов разный: где-то простое поле ean13, где-то productCode,
    где-то вложенный productCodeNew. Марочные коды (КИЗ) сюда не годятся — из них
    берём только часть, похожую на EAN, и только если она отдельным полем.
    """
    for key in ("ean13", "barcode", "productCode", "rawProductCode", "gtin", "ean"):
        value = item.get(key) or item.get(key.capitalize())
        if isinstance(value, (str, int)):
            match = _EAN.search(str(value))
            if match:
                return match.group(1)
    nested = item.get("productCodeNew") or item.get("ProductCodeNew")
    if isinstance(nested, dict):
        for node in nested.values():
            if isinstance(node, dict):
                found = _barcode(node)
                if found:
                    return found
    return None


def _find_receipt_node(node: Any) -> dict | None:
    """Первый объект со списком items — у разных ОФД он лежит на разной глубине."""
    if isinstance(node, dict):
        if isinstance(node.get("items"), list) and node["items"]:
            return node
        for value in node.values():
            found = _find_receipt_node(value)
            if found:
                return found
    elif isinstance(node, list):
        for value in node:
            found = _find_receipt_node(value)
            if found:
                return found
    return None


def parse_receipt_json(payload: str | dict) -> Receipt:
    data = json.loads(payload) if isinstance(payload, str) else payload
    node = _find_receipt_node(data)
    if not node:
        raise ValueError("В файле не нашёлся список позиций (items)")

    rows: list[ReceiptRow] = []
    for item in node["items"]:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or item.get("Name") or "").strip()
        if not name:
            continue
        qty = _number(item.get("quantity", item.get("Quantity", 1))) or 1.0
        price = _money(item.get("price", item.get("Price")))
        total = _money(item.get("sum", item.get("Sum")))
        if price is None and total is not None:
            price = round(total / qty, 2) if qty else total
        if total is None and price is not None:
            total = round(price * qty, 2)
        if price is None or total is None:
            continue
        rows.append(ReceiptRow(raw_name=name, qty=qty, unit_price=price, total=total,
                               barcode=_barcode(item)))

    if not rows:
        raise ValueError("В файле нет ни одной позиции с ценой")

    store = (node.get("retailPlace") or node.get("user") or node.get("store")
             or node.get("retailPlaceAddress") or "—")
    total = _money(node.get("totalSum", node.get("total")))
    return Receipt(
        date=_date(node.get("dateTime") or node.get("date")),
        store_name=str(store).strip() or "—",
        rows=rows,
        total=total if total is not None else round(sum(r.total for r in rows), 2),
    )


# ---------- таблица ----------
def _rows_from_csv(path: str) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            with open(path, encoding=encoding, newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                except csv.Error:
                    dialect = csv.excel
                return [row for row in csv.reader(fh, dialect) if any(str(c).strip() for c in row)]
        except UnicodeDecodeError:
            continue
    raise ValueError("Не удалось прочитать файл — проверьте кодировку")


def _rows_from_xlsx(path: str) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("Для XLSX нужна библиотека openpyxl") from exc
    sheet = load_workbook(path, data_only=True).active
    return [list(row) for row in sheet.iter_rows(values_only=True) if any(c is not None and str(c).strip() for c in row)]


def parse_receipt_table(path: str) -> Receipt:
    ext = os.path.splitext(path)[1].lower()
    table = _rows_from_xlsx(path) if ext in (".xlsx", ".xlsm") else _rows_from_csv(path)
    if len(table) < 2:
        raise ValueError("В таблице нет строк с данными")

    header = [str(c or "") for c in table[0]]
    idx = {kind: _pick_column(header, kind) for kind in COLUMNS}
    if idx["name"] is None:
        raise ValueError("Не нашлась колонка с названием товара")
    if idx["price"] is None and idx["total"] is None:
        raise ValueError("Нужна хотя бы одна колонка с ценой или суммой")

    def cell(row: list[Any], kind: str) -> Any:
        i = idx[kind]
        return row[i] if i is not None and i < len(row) else None

    rows: list[ReceiptRow] = []
    date_value = store_value = None
    for row in table[1:]:
        name = str(cell(row, "name") or "").strip()
        if not name:
            continue
        qty = _number(cell(row, "qty"))
        qty = qty if qty and qty > 0 else 1.0
        price = _number(cell(row, "price"))
        total = _number(cell(row, "total"))
        if price is None and total is not None:
            price = round(total / qty, 2)
        if total is None and price is not None:
            total = round(price * qty, 2)
        if price is None or total is None:
            continue
        rows.append(ReceiptRow(raw_name=name, qty=qty, unit_price=round(price, 2), total=round(total, 2)))
        date_value = date_value or cell(row, "date")
        store_value = store_value or cell(row, "store")

    if not rows:
        raise ValueError("В таблице не нашлось ни одной позиции с ценой")

    return Receipt(
        date=_date(date_value),
        store_name=str(store_value or "—").strip() or "—",
        rows=rows,
        total=round(sum(r.total for r in rows), 2),
    )
